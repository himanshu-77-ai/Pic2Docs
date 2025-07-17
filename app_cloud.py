import streamlit as st
import easyocr
import numpy as np
from PIL import Image, ImageEnhance
from io import BytesIO
from fpdf import FPDF
from docx import Document
import pandas as pd
import json
import time

# ------------------------ OCR SETUP ------------------------
@st.cache_resource
def load_ocr_model(languages=['en']):
    """Load OCR model only for required languages to save memory"""
    try:
        return easyocr.Reader(languages, gpu=False)
    except Exception as e:
        st.error(f"Model loading failed: {str(e)}")
        return None

def preprocess_image(image):
    """Enhance image for better OCR accuracy"""
    img = image.convert('L')  # Convert to grayscale
    enhancer = ImageEnhance.Contrast(img)
    img = enhancer.enhance(2.0)  # Increase contrast
    return np.array(img)

def ocr_image_local(image_file, language_code="en"):
    """Extract text from image with language support"""
    try:
        reader = load_ocr_model([language_code])
        if reader is None:
            return "❌ Model not loaded"
        
        image = Image.open(image_file)
        processed_img = preprocess_image(image)
        
        start_time = time.time()
        result = reader.readtext(processed_img, detail=0, paragraph=True)
        processing_time = round(time.time() - start_time, 2)
        
        st.success(f"Processed in {processing_time}s | Language: {language_code}")
        return '\n'.join(result) if result else "❌ No text found"
    except Exception as e:
        return f"❌ OCR Failed: {str(e)}"

# ------------------------ EXPORT FUNCTIONS ------------------------
def convert_to_pdf(text):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    for line in text.split('\n'):
        pdf.cell(200, 10, txt=line, ln=True)
    return pdf.output(dest='S').encode('latin-1')

def convert_to_word(text):
    doc = Document()
    doc.add_paragraph(text)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()

def convert_to_excel(text):
    lines = text.split('\n')
    df = pd.DataFrame({"Extracted Text": lines})
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine='xlsxwriter') as writer:
        df.to_excel(writer, index=False)
    return buffer.getvalue()

def convert_to_ipynb(text):
    notebook = {
        "cells": [{
            "cell_type": "markdown",
            "metadata": {},
            "source": [text]
        }],
        "metadata": {},
        "nbformat": 4,
        "nbformat_minor": 2
    }
    return json.dumps(notebook, indent=2).encode('utf-8')

# ------------------------ STREAMLIT UI ------------------------
st.set_page_config(page_title="📸 Pic2Docs", layout="centered")
st.title("📸 Pic2Docs - Image to Document Converter")
st.markdown("""
Extract text from images in multiple languages and export to various formats.
**Works offline** - No API keys needed ✅
""")

# Language mapping
LANG_MAP = {
    "English": "en",
    "Hindi": "hi",
    "French": "fr",
    "Spanish": "es",
    "German": "de",
    "Bengali": "bn"
}

# File upload section
uploaded_file = st.file_uploader("Upload an image", type=["png", "jpg", "jpeg"])
if uploaded_file:
    col1, col2 = st.columns(2)
    with col1:
        st.image(uploaded_file, caption="Uploaded Image", use_column_width=True)
    with col2:
        lang_display = st.selectbox("Select text language", list(LANG_MAP.keys()))
        if lang_display == "Bengali":
            st.warning("For best results with Bengali:")
            st.markdown("- Use high-contrast images\n- Clear printed text works best")

    if st.button("🔍 Extract Text", type="primary"):
        with st.spinner(f"Processing {lang_display} text..."):
            extracted_text = ocr_image_local(uploaded_file, LANG_MAP[lang_display])

        st.subheader("📄 Extracted Text")
        st.text_area("", extracted_text, height=300, label_visibility="collapsed")

        if not extracted_text.startswith("❌"):
            st.download_button("📥 Download as TXT", extracted_text, file_name="extracted_text.txt")
            st.download_button("📄 Download as PDF", convert_to_pdf(extracted_text), file_name="extracted_text.pdf")
            st.download_button("📝 Download as Word", convert_to_word(extracted_text), file_name="extracted_text.docx")
            st.download_button("📊 Download as Excel", convert_to_excel(extracted_text), file_name="extracted_text.xlsx")
else:
    st.info("👆 Please upload an image file to get started")

# Footer
st.markdown("---")
st.caption("Made with ♥ by Pic2Docs | Supports English, Hindi, Bengali, French, Spanish & German")
